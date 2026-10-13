"""Pure report-composition/render tests; no MongoDB test double is required."""
from io import BytesIO
import re

from pptx import Presentation
from pypdf import PdfReader

from backend.app.reporting import make_report_model, generate_pdf, generate_powerpoint


def _case():
    return {
        "_id": "TRACEX-REPORT-TEST",
        "case_information": {"case_name": "Controlled report test"},
        "summary": {"synthetic": True},
        "status": "Investigating",
    }


def test_model_preserves_timestamp_values_sorts_chronologically_and_labels_relationship_basis():
    events = [
        {"event_id": "later", "timestamp": "2026-10-11T10:07:00Z", "event_type": "token_reuse"},
        {"event_id": "earlier", "timestamp": "2026-10-11T10:05:00Z", "event_type": "login"},
    ]
    analysis = {"correlation": [{"event_ids": ["earlier", "later"], "link_reasons": ["shared session"]}]}
    model = make_report_model(
        _case(), analysis, events, [],
        {"nodes": ["identity-a", "service-a"], "edges": [{"from": "identity-a", "to": "service-a", "permission": "read"}]},
    )

    assert [event["event_id"] for event in model["events"]] == ["earlier", "later"]
    assert [event["timestamp"] for event in model["events"]] == ["2026-10-11T10:05:00Z", "2026-10-11T10:07:00Z"]
    assert {edge["basis"] for edge in model["relationships"]} == {
        "Explicit environment relationship", "Inferred by existing event correlation"
    }
    assert model["root_cause_text"] == "No root-cause analysis is available in the persisted investigation data."


def test_pdf_and_editable_powerpoint_render_and_parse_with_source_disclosures():
    model = make_report_model(
        _case(), None,
        [{"event_id": "event-a", "timestamp": "2026-10-11T10:05:00Z", "event_type": "login", "summary": "Observed sign-in."}],
        [], None,
    )
    sections = ["executive_summary", "timeline", "relationship_map", "root_cause", "methodology"]

    pdf_bytes = generate_pdf(model, sections)
    pages = PdfReader(BytesIO(pdf_bytes), strict=True).pages
    pdf_text = "\n".join(page.extract_text() or "" for page in pages)
    assert len(pages) >= 2
    assert "Controlled report test" in pdf_text
    assert "2026-10-11T10:05:00Z" in re.sub(r"\s+", "", pdf_text)
    assert "No AI provider is configured" in pdf_text

    deck_bytes = generate_powerpoint(model, sections)
    deck = Presentation(BytesIO(deck_bytes))
    slide_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert len(deck.slides) >= 4
    assert "Controlled report test" in slide_text
    assert "No AI provider is configured" in slide_text
