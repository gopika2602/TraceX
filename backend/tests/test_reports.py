"""Report generation tests use isolated synthetic investigations and Mongo-compatible storage."""
from uuid import uuid4

import mongomock
import mongomock.gridfs
import pytest
from fastapi.testclient import TestClient
from gridfs import GridFSBucket
from pptx import Presentation
from pypdf import PdfReader
from io import BytesIO

from backend.app.db import ensure_indexes, get_database
from backend.app.main import app
from backend.tests.test_api_workflow import register


@pytest.fixture
def api(monkeypatch):
    mongomock.gridfs.enable_gridfs_integration()
    database = mongomock.MongoClient(tz_aware=True)[f"tracex-report-test-{uuid4().hex}"]
    ensure_indexes(database)
    from backend.app import main as main_module
    monkeypatch.setattr(main_module, "get_database", lambda: database)
    monkeypatch.setattr(main_module, "close_client", lambda: None)
    monkeypatch.delenv("ADMIN_EMAIL", raising=False)
    monkeypatch.delenv("ADMIN_PASSWORD", raising=False)
    app.dependency_overrides[get_database] = lambda: database
    try:
        with TestClient(app) as client:
            yield client, database
    finally:
        app.dependency_overrides.clear()


def create_case(api: TestClient, *, with_evidence: bool = True):
    if with_evidence:
        response = api.post("/cases/intake", json={
            "case_name": "Controlled report test",
            "description": "Synthetic records for regression testing.",
            "evidence": [
                {"evidence_type": "login", "timestamp": "2026-10-11T10:05:00Z", "identity": "analyst@example.test", "device_id": "device-a", "source_ip": "203.0.113.9", "description": "Suspicious sign-in observed.", "verification_status": "approved", "raw_event_reference": "test:row:1"},
                {"evidence_type": "token_reuse", "timestamp": "2026-10-11T10:07:00Z", "identity": "analyst@example.test", "session_id": "session-a", "description": "Session was reused.", "verification_status": "unreviewed", "raw_event_reference": "test:row:2"},
            ],
        })
        assert response.status_code == 201, response.text
        return response.json()["case_id"]
    response = api.post("/datasets/load-demo")
    assert response.status_code == 201, response.text
    created = api.post("/cases", json={"dataset_id": response.json()["id"]})
    assert created.status_code == 201, created.text
    return created.json()["case_id"]


def test_report_readiness_is_authenticated_owner_scoped_and_counts_persisted_records(api):
    client, _ = api
    assert client.get("/cases/anything/reports/readiness").status_code == 401
    assert register(client, "report-owner@example.test").status_code == 201
    case_id = create_case(client)
    ready = client.get(f"/cases/{case_id}/reports/readiness")
    assert ready.status_code == 200, ready.text
    body = ready.json()
    assert body["counts"] == {"evidence_records": 2, "approved": 1, "rejected": 0, "awaiting_review": 1, "timeline_events": 2}
    assert body["analysis_available"] is False
    assert body["composition_mode"] == "deterministic"
    assert body["sections"][0]["available"] is True

    client.post("/auth/logout")
    assert register(client, "other-user@example.test").status_code == 201
    assert client.get(f"/cases/{case_id}/reports/readiness").status_code == 404
    assert client.get(f"/cases/{case_id}/reports").status_code == 404


def test_generate_both_validates_files_persists_metadata_and_downloads(api):
    client, database = api
    assert register(client, "reporter@example.test").status_code == 201
    case_id = create_case(client)
    analysis = client.post(f"/cases/{case_id}/analyze")
    assert analysis.status_code == 200, analysis.text
    sections = ["executive_summary", "evidence_summary", "timeline", "relationship_map", "attack_path", "root_cause", "unresolved_questions", "methodology"]
    generated = client.post(f"/cases/{case_id}/reports", json={"format": "both", "sections": sections})
    assert generated.status_code == 201, generated.text
    report = generated.json()["report"]
    assert report["status"] == "ready"
    assert {item["format"] for item in report["files"]} == {"pdf", "powerpoint"}
    assert all(item["size_bytes"] > 0 and len(item["sha512"]) == 128 for item in report["files"])

    pdf = client.get(f"/reports/{report['id']}/pdf/download")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")
    parsed_pdf = PdfReader(BytesIO(pdf.content), strict=True)
    pdf_text = "\n".join(page.extract_text() or "" for page in parsed_pdf.pages)
    assert "Controlled report test" in pdf_text
    assert "test:row:1" in pdf_text
    assert "Incident timeline" in pdf_text
    assert "relationship" in pdf_text.lower()

    pptx = client.get(f"/reports/{report['id']}/powerpoint/download")
    assert pptx.status_code == 200
    deck = Presentation(BytesIO(pptx.content))
    assert len(deck.slides) >= 5
    slide_text = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "Controlled report test" in slide_text
    assert "Incident Timeline" in slide_text
    assert any(shape.shape_type == 1 for slide in deck.slides for shape in slide.shapes)

    metadata = database.reports.find_one({"_id": report["id"]})
    assert metadata["source_evidence_ids"]
    assert len(metadata["source_event_ids"]) == 2
    assert len(list(GridFSBucket(database, bucket_name="tracex_reports").find({"metadata.report_id": report["id"]}))) == 2
    listing = client.get(f"/cases/{case_id}/reports")
    assert listing.status_code == 200 and len(listing.json()["reports"]) == 1


def test_report_with_no_evidence_and_incomplete_analysis_is_honest(api):
    client, _ = api
    assert register(client, "incomplete@example.test").status_code == 201
    case_id = create_case(client, with_evidence=False)
    ready = client.get(f"/cases/{case_id}/reports/readiness").json()
    assert ready["counts"]["evidence_records"] == 0
    assert ready["counts"]["timeline_events"] > 0
    request = client.post(f"/cases/{case_id}/reports", json={"format": "pdf", "sections": ["executive_summary", "evidence_summary", "unresolved_questions"]})
    assert request.status_code == 201, request.text
    pdf = client.get(request.json()["report"]["files"][0]["download_url"])
    text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf.content)).pages)
    assert "No persisted evidence records" in text or "0 evidence" in text
    assert "No AI provider is configured" in text


def test_report_rejects_invalid_sections_and_blocks_unauthorized_download(api):
    client, _ = api
    assert register(client, "report-security@example.test").status_code == 201
    case_id = create_case(client)
    assert client.post(f"/cases/{case_id}/reports", json={"format": "pdf", "sections": ["made_up"]}).status_code == 422
    report_result = client.post(f"/cases/{case_id}/reports", json={"format": "pdf", "sections": ["executive_summary"]})
    assert report_result.status_code == 201
    report_id = report_result.json()["report"]["id"]
    client.post("/auth/logout")
    assert register(client, "unauthorized@example.test").status_code == 201
    assert client.get(f"/reports/{report_id}/pdf/download").status_code == 404
    assert client.get(f"/reports/{report_id}/invalid/download").status_code == 422
