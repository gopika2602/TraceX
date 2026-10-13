"""Evidence-grounded report model and PDF/PowerPoint renderers.

This module only formats persisted TraceX case records and engine output. It does
not call an AI provider, change investigation data, or infer new conclusions.
"""
from __future__ import annotations

from datetime import datetime, timezone
from html import escape
from io import BytesIO
import math
from typing import Any
from pydantic import BaseModel, ConfigDict, Field, model_validator

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, Frame, PageBreak, PageTemplate,
    Paragraph, KeepTogether, Spacer, Table, TableStyle, Flowable)
from reportlab.graphics.shapes import Drawing, Line, Circle, String, Rect
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pypdf import PdfReader


SECTION_LABELS = {
    "executive_summary": "Executive summary", "evidence_summary": "Evidence inventory and review status",
    "timeline": "Incident timeline", "relationship_map": "Entity relationship map",
    "attack_path": "Potential attack path", "root_cause": "Root-cause analysis",
    "risk_assessment": "Risk assessment and potential impact", "remediation": "Recommended remediation",
    "unresolved_questions": "Unresolved questions and evidence gaps",
    "technical_appendix": "Technical appendix", "submitted_images": "Submitted evidence images",
    "methodology": "Investigation methodology",
}
MAX_TIMELINE_EVENTS = 2000
MAX_REPORT_BYTES = 25 * 1024 * 1024
NAVY = colors.HexColor("#101828")
PINK = colors.HexColor("#E86ABF")
CYAN = colors.HexColor("#35D8FF")
INK = colors.HexColor("#263347")
MUTED = colors.HexColor("#667085")
PALE = colors.HexColor("#F4F2F8")


class ValidatedReportModel(BaseModel):
    """Minimal schema contract for source-grounded report composition."""
    model_config = ConfigDict(extra="forbid")
    case_id: str = Field(min_length=1)
    title: str = Field(min_length=1)
    status: str
    created_at: Any = None
    detected: Any = None
    organization: Any = None
    incident_at: Any = None
    description: Any = None
    investigator_name: Any = None
    investigator_notes: Any = None
    generated_at: str
    composition_mode: str
    composition_disclosure: str
    evidence: list[dict[str, Any]]
    evidence_counts: dict[str, int]
    events: list[dict[str, Any]]
    event_count: int
    events_in_report: int
    relationships: list[dict[str, Any]]
    relationship_nodes: list[dict[str, Any]]
    attack_path: list[dict[str, Any]]
    root_cause: dict[str, Any]
    root_cause_text: str
    root_cause_status: str
    blast_radius: dict[str, Any]
    remediations: list[dict[str, Any]]
    attack_origin: dict[str, Any]
    limitations: list[str]
    available_sections: dict[str, bool]
    analysis_available: bool
    synthetic: bool

    @model_validator(mode="after")
    def validate_counts(self):
        if self.events_in_report > self.event_count:
            raise ValueError("Report timeline cannot claim more events than the saved investigation.")
        if self.evidence_counts.get("total", -1) != len(self.evidence):
            raise ValueError("Report evidence count does not match the retrieved evidence records.")
        return self


def _clean(value: Any, fallback: str = "Unknown") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list)):
        return str(value)
    return str(value)


def _timestamp(event: dict[str, Any]) -> str:
    return _clean(event.get("timestamp", event.get("time", event.get("occurred_at"))), "Time not recorded")


def _event_id(event: dict[str, Any], index: int) -> str:
    return _clean(event.get("event_id", event.get("id")), f"record-{index + 1}")


def make_report_model(case: dict[str, Any], analysis: dict[str, Any] | None,
                      events: list[dict[str, Any]], evidence_rows: list[dict[str, Any]],
                      environment: dict[str, Any] | None, *, generated_at: datetime | None = None) -> dict[str, Any]:
    """Build a bounded, JSON-safe report model from already-authorized persisted records."""
    summary = case.get("summary", {})
    info = case.get("case_information", {})
    result = analysis or {}
    def event_sort_key(event: dict[str, Any]):
        raw = event.get("timestamp", event.get("time", event.get("occurred_at")))
        try:
            value = datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
            value = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)
            return (False, value)
        except (TypeError, ValueError):
            return (True, datetime.max.replace(tzinfo=timezone.utc))
    event_rows = sorted(events, key=event_sort_key)
    event_rows = event_rows[:MAX_TIMELINE_EVENTS]
    evidence: list[dict[str, Any]] = []
    for row in evidence_rows:
        record = row.get("record", {})
        evidence.append({"id": str(row.get("_id", "")), "event_id": str(row.get("event_id", "")),
            "evidence_type": _clean(record.get("evidence_type")), "timestamp": _clean(record.get("timestamp"), "Time not recorded"),
            "description": _clean(record.get("description"), "No description recorded"),
            "source_reference": _clean(record.get("raw_event_reference", record.get("evidence_source", record.get("source"))), "Not recorded"),
            "review_status": _clean(record.get("verification_status"), "awaiting review"),
            "integrity_hash": record.get("integrity_hash") or record.get("sha256") or None,
        })
    count_status = {name: sum(1 for item in evidence if item["review_status"].lower() == name)
                    for name in ("approved", "rejected", "unreviewed")}
    count_status["awaiting_review"] = sum(1 for item in evidence if item["review_status"].lower() not in {"approved", "rejected"})
    event_ids = {_event_id(item, i) for i, item in enumerate(event_rows)}

    edges: list[dict[str, Any]] = []
    nodes: dict[str, dict[str, str]] = {}
    for node in (environment or {}).get("nodes", []):
        node_id = node if isinstance(node, str) else node.get("id", node.get("node_id")) if isinstance(node, dict) else None
        if node_id:
            attrs = {} if isinstance(node, str) else node
            nodes[str(node_id)] = {"id": str(node_id), "type": _clean(attrs.get("type"), "entity"), "basis": "Recorded environment data"}
    for edge in (environment or {}).get("edges", []):
        source = edge.get("from", edge.get("source")); target = edge.get("to", edge.get("target"))
        if source is not None and target is not None:
            source, target = str(source), str(target)
            edges.append({"source": source, "target": target,
                          "label": _clean(edge.get("permission", edge.get("permissions", edge.get("actions"))), "recorded access"),
                          "basis": "Explicit environment relationship"})
            nodes.setdefault(source, {"id": source, "type": "entity", "basis": "Recorded environment data"})
            nodes.setdefault(target, {"id": target, "type": "entity", "basis": "Recorded environment data"})
    for group in result.get("correlation", []) if isinstance(result.get("correlation", []), list) else []:
        ids = [str(item) for item in group.get("event_ids", []) if str(item) in event_ids]
        reasons = ", ".join(map(str, group.get("link_reasons", []))) or "shared evidence identifiers"
        for first, second in zip(ids, ids[1:]):
            edges.append({"source": first, "target": second, "label": reasons, "basis": "Inferred by existing event correlation"})
            nodes.setdefault(first, {"id": first, "type": "event", "basis": "Persisted event"})
            nodes.setdefault(second, {"id": second, "type": "event", "basis": "Persisted event"})

    path = result.get("attack_path", {}) if isinstance(result.get("attack_path", {}), dict) else {}
    steps = path.get("steps", []) if isinstance(path.get("steps", []), list) else []
    root = result.get("root_cause", {}) if isinstance(result.get("root_cause", {}), dict) else {}
    blast = result.get("blast_radius", {}) if isinstance(result.get("blast_radius", {}), dict) else {}
    remediation = result.get("remediations", []) if isinstance(result.get("remediations", []), list) else []
    origin = result.get("attack_origin", {}) if isinstance(result.get("attack_origin", {}), dict) else {}
    limitations = [str(item) for item in root.get("limitations", [])]
    if not analysis:
        limitations.append("TraceX analysis is not available for this investigation.")
    if not events:
        limitations.append("No persisted timeline events are associated with this investigation.")
    if not environment:
        limitations.append("No explicit environment graph was available for this report.")
    observed_gaps = []
    for i, event in enumerate(event_rows):
        missing = [name for name, value in (("source IP", (event.get("source") or {}).get("ip") if isinstance(event.get("source"), dict) else event.get("source_ip")),
                    ("device", event.get("device_id")), ("identity", event.get("identity"))) if not value]
        if missing:
            observed_gaps.append(f"{_event_id(event, i)}: missing {', '.join(missing)}.")
    limitations.extend(observed_gaps[:100])

    available = {
        "executive_summary": True,
        "evidence_summary": bool(evidence),
        "timeline": bool(events),
        "relationship_map": bool(edges),
        "attack_path": bool(steps),
        "root_cause": bool(root),
        "risk_assessment": bool(blast or any(event.get("severity") for event in event_rows)),
        "remediation": bool(remediation),
        "unresolved_questions": bool(limitations or origin),
        "technical_appendix": bool(events or evidence),
        "submitted_images": False,
        "methodology": True,
    }
    cause_text = _clean(root.get("summary"), "The available analysis did not establish a root cause.")
    if not root:
        cause_text = "No root-cause analysis is available in the persisted investigation data."
    status = "supported by existing analysis" if root else "unknown"
    if root and root.get("confidence") is not None:
        status = "supported by existing analysis; confidence is reported by the engine"
    exec_summary = cause_text if root else f"Investigation {case.get('_id', '')} has no persisted root-cause finding. See the saved evidence and data gaps below."
    if result.get("synthetic") or summary.get("synthetic"):
        limitations.insert(0, "This investigation is marked as a controlled synthetic dataset.")

    model = {
        "case_id": str(case.get("_id", "")), "title": _clean(info.get("case_name", summary.get("title")), "Untitled investigation"),
        "status": _clean(case.get("status", summary.get("status"))), "created_at": _clean(case.get("created_at")),
        "detected": _clean(summary.get("detected")), "organization": info.get("organization"),
        "incident_at": info.get("incident_at"), "description": info.get("description"),
        "investigator_name": info.get("investigator_name"), "investigator_notes": info.get("investigator_notes"),
        "generated_at": (generated_at or datetime.now(timezone.utc)).astimezone(timezone.utc).isoformat(),
        "composition_mode": "deterministic", "composition_disclosure": "No AI provider is configured. Text is prepared deterministically from saved case records and existing analysis; the report does not independently validate the underlying claims.",
        "evidence": evidence, "evidence_counts": {"total": len(evidence), **count_status},
        "events": event_rows, "event_count": len(events), "events_in_report": len(event_rows),
        "relationships": edges, "relationship_nodes": list(nodes.values()), "attack_path": steps,
        "root_cause": root, "root_cause_text": cause_text, "root_cause_status": status,
        "blast_radius": blast, "remediations": remediation, "attack_origin": origin,
        "limitations": list(dict.fromkeys(limitations)), "available_sections": available,
        "analysis_available": bool(analysis), "synthetic": bool(result.get("synthetic") or summary.get("synthetic")),
    }
    return ValidatedReportModel.model_validate(model).model_dump()


def _safe(value: Any, fallback: str = "Not recorded") -> str:
    return escape(_clean(value, fallback)).replace("\n", "<br/>")


def _paragraph(text: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_safe(text), style)


def _relationship_drawing(model: dict[str, Any], width: float = 460, height: float = 250) -> Drawing:
    drawing = Drawing(width, height)
    nodes = model["relationship_nodes"][:24]
    edges = model["relationships"]
    if not nodes or not edges:
        drawing.add(String(12, height / 2, "No recorded or correlated relationships are available.", fontSize=10, fillColor=MUTED))
        return drawing
    by_id = {item["id"]: item for item in nodes}
    cx, cy = width / 2, height / 2
    radius = min(width, height) * .36
    positions = {node["id"]: (cx + radius * math.cos(2 * math.pi * i / len(nodes)), cy + radius * math.sin(2 * math.pi * i / len(nodes))) for i, node in enumerate(nodes)}
    for edge in edges[:40]:
        if edge["source"] not in positions or edge["target"] not in positions:
            continue
        x1, y1 = positions[edge["source"]]; x2, y2 = positions[edge["target"]]
        stroke = CYAN if edge["basis"] == "Explicit environment relationship" else PINK
        drawing.add(Line(x1, y1, x2, y2, strokeColor=stroke, strokeWidth=1))
    for node in nodes:
        x, y = positions[node["id"]]
        drawing.add(Circle(x, y, 4.5, fillColor=CYAN if node["type"] != "event" else PINK, strokeColor=NAVY))
        label = node["id"][:22]
        drawing.add(String(x + 7, y - 3, label, fontSize=6.5, fillColor=INK))
    drawing.add(String(4, 7, "Cyan: explicit environment edge  ·  Pink: existing-analysis correlation (inferred)", fontSize=7, fillColor=MUTED))
    return drawing


class ReportCanvasDoc(BaseDocTemplate):
    def __init__(self, stream: BytesIO, **kwargs: Any):
        super().__init__(stream, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm, topMargin=20*mm, bottomMargin=18*mm, **kwargs)
        frame = Frame(self.leftMargin, self.bottomMargin, self.width, self.height, id="normal")
        self.addPageTemplates([PageTemplate(id="tracex", frames=frame, onPage=self._decorate)])

    def _decorate(self, canvas, doc):
        canvas.saveState()
        if doc.page > 1:
            canvas.setFillColor(NAVY); canvas.rect(0, A4[1]-13*mm, A4[0], 13*mm, stroke=0, fill=1)
            canvas.setFillColor(colors.white); canvas.setFont("Helvetica-Bold", 9); canvas.drawString(18*mm, A4[1]-8.5*mm, "TraceX  /  Investigation Report")
            canvas.setFillColor(MUTED); canvas.setFont("Helvetica", 8); canvas.drawCentredString(A4[0]/2, 9*mm, f"{doc.page}")
            canvas.setStrokeColor(PINK); canvas.setLineWidth(1); canvas.line(18*mm, 14*mm, A4[0]-18*mm, 14*mm)
        canvas.restoreState()


def generate_pdf(model: dict[str, Any], sections: list[str]) -> bytes:
    output = BytesIO()
    doc = ReportCanvasDoc(output, title=f"TraceX Report — {model['title']}", author="TraceX")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TXTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=29, leading=34, textColor=NAVY, alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="TXHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=NAVY, spaceBefore=7, spaceAfter=9, keepWithNext=True))
    styles.add(ParagraphStyle(name="TXBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=13, textColor=INK, spaceAfter=7, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="TXSmall", parent=styles["BodyText"], fontSize=7.5, leading=10, textColor=MUTED, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="TXCenter", parent=styles["TXBody"], alignment=TA_CENTER))
    story: list[Flowable] = [Spacer(1, 45*mm), Paragraph("Trace<span color='#E86ABF'>X</span>", styles["TXTitle"]),
        Paragraph("INVESTIGATION REPORT", styles["TXCenter"]), Spacer(1, 15*mm),
        Paragraph(_safe(model["title"]), styles["TXTitle"]),
        Paragraph(f"Investigation ID: {_safe(model['case_id'])}", styles["TXCenter"]),
        Paragraph(f"Status: {_safe(model['status'])}", styles["TXCenter"]), Spacer(1, 8*mm),
        Paragraph(f"Generated: {_safe(model['generated_at'])}", styles["TXCenter"]),
        Paragraph("Report composition: deterministic, based on persisted TraceX data", styles["TXCenter"])]
    if model["synthetic"]:
        story += [Spacer(1, 8*mm), Paragraph("CONTROLLED SYNTHETIC DATASET", styles["TXCenter"])]
    story += [PageBreak(), Paragraph("Investigation Overview", styles["TXHeading"]),
        Paragraph(_safe(model["composition_disclosure"]), styles["TXBody"]),
        Paragraph(f"Detected: {_safe(model['detected'])}  ·  Events: {model['event_count']}  ·  Saved evidence records: {model['evidence_counts']['total']}", styles["TXBody"])]
    for key in sections:
        if key == "executive_summary":
            story += [Paragraph(SECTION_LABELS[key], styles["TXHeading"]), Paragraph(_safe(model["root_cause_text"]), styles["TXBody"]), Paragraph(f"Assessment status: {_safe(model['root_cause_status'])}", styles["TXSmall"])]
        elif key == "evidence_summary":
            story.append(Paragraph(SECTION_LABELS[key], styles["TXHeading"]))
            rows = [["Evidence ID", "Type", "Timestamp", "Review", "Source reference", "Description"]]
            rows.extend([[item["id"], item["evidence_type"], item["timestamp"], item["review_status"], item["source_reference"], item["description"]] for item in model["evidence"]])
            if len(rows) == 1:
                rows.append(["No persisted evidence records", "—", "—", "—", "—", "—"])
            widths = [40, 50, 69, 47, 80, doc.width-286]
            table = Table([[ _paragraph(cell, styles["TXSmall"]) for cell in row] for row in rows], colWidths=widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#D5D9E2")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,PALE]),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
            story.append(table)
            story.append(Paragraph(f"Review counts: {model['evidence_counts']['approved']} approved, {model['evidence_counts']['rejected']} rejected, {model['evidence_counts']['awaiting_review']} awaiting review. Approval records a review decision; it does not independently verify the underlying claim.", styles["TXBody"]))
        elif key == "timeline":
            story.append(Paragraph(SECTION_LABELS[key], styles["TXHeading"]))
            rows = [["Timestamp (source value)", "Event", "Identity", "Evidence ref", "Description"]]
            rows.extend([[_timestamp(item), _clean(item.get("event_type", item.get("type"))), _clean(item.get("identity")), _event_id(item, index), _clean(item.get("summary", item.get("description")), "No description")]
                         for index, item in enumerate(model["events"])])
            if len(rows) == 1: rows.append(["No events", "—", "—", "—", "—"])
            table=Table([[_paragraph(cell,styles["TXSmall"]) for cell in row] for row in rows],colWidths=[74,68,70,68,doc.width-280],repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D5D9E2")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,PALE]),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
            story += [table, Paragraph(f"Source event count: {model['event_count']}; included: {model['events_in_report']}. Source timestamp strings are preserved; missing timestamps are labeled.",styles["TXSmall"])]
        elif key == "relationship_map":
            story += [Paragraph(SECTION_LABELS[key],styles["TXHeading"]), _relationship_drawing(model),
                Paragraph("Explicit environment relationships and engine-inferred event correlations use different labels and colors. Correlation is an analytical relationship, not proof of causation.",styles["TXSmall"])]
            if model["relationships"]:
                rows=[["From","To","Relationship","Basis"]]+[[edge["source"],edge["target"],edge["label"],edge["basis"]] for edge in model["relationships"]]
                table=Table([[_paragraph(cell,styles["TXSmall"]) for cell in row] for row in rows],colWidths=[80,80,doc.width-250,90],repeatRows=1)
                table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D5D9E2")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,PALE])]))
                story.append(table)
        elif key == "attack_path":
            story.append(Paragraph(SECTION_LABELS[key],styles["TXHeading"]))
            if model["attack_path"]:
                rows=[["Step","Timestamp","Event / transition","Identity","Evidence IDs","Assessment"]]
                for i, step in enumerate(model["attack_path"],1):
                    rows.append([i,step.get("timestamp") or "Time not recorded",step.get("event_type",step.get("action","Observed step")),step.get("identity","Unknown"),", ".join(map(str,step.get("event_ids",[]))),step.get("status","Existing engine output")])
                story.append(Table([[_paragraph(c,styles["TXSmall"]) for c in row] for row in rows],colWidths=[30,70,95,65,90,doc.width-350],repeatRows=1,style=TableStyle([("BACKGROUND",(0,0),(-1,0),NAVY),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.3,colors.HexColor("#D5D9E2")),("VALIGN",(0,0),(-1,-1),"TOP"),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,PALE])])) )
                story.append(Paragraph("This is the existing engine's analysis output; transitions are not independently validated by report generation.",styles["TXSmall"]))
            else: story.append(Paragraph("No supported attack path is available in the persisted analysis.",styles["TXBody"]))
        elif key == "root_cause":
            story += [Paragraph(SECTION_LABELS[key],styles["TXHeading"]),Paragraph(_safe(model["root_cause_text"]),styles["TXBody"]),Paragraph(f"Status: {_safe(model['root_cause_status'])}",styles["TXBody"])]
            for factor in model["root_cause"].get("contributing_factors",[]): story.append(Paragraph(f"• {_safe(factor.get('code'))}: {_safe(factor.get('description'))}",styles["TXBody"]))
        elif key == "risk_assessment":
            story.append(Paragraph(SECTION_LABELS[key],styles["TXHeading"]))
            blast=model["blast_radius"]
            if blast:
                story.append(Paragraph("Existing engine blast-radius output (not an independent report-agent calculation):",styles["TXBody"]))
                story.append(Paragraph(_safe(blast.get("summary",{})),styles["TXBody"]))
                observed=blast.get("observed_targets",[])
                if observed: story.append(Paragraph("Observed destination targets in event telemetry: " + _safe([item.get("node_id") for item in observed]),styles["TXBody"]))
            else: story.append(Paragraph("No quantitative risk score or blast-radius result is available. This report does not invent one.",styles["TXBody"]))
        elif key == "remediation":
            story.append(Paragraph(SECTION_LABELS[key],styles["TXHeading"]))
            if model["remediations"]:
                for item in model["remediations"]: story.append(Paragraph(f"• {_safe(item.get('title',item.get('action')))} — {_safe(item.get('description',item.get('reason')))} [{_safe(item.get('status','recommended'))}]",styles["TXBody"]))
            else: story.append(Paragraph("No remediation recommendations are saved for this analysis.",styles["TXBody"]))
        elif key == "unresolved_questions":
            story.append(Paragraph(SECTION_LABELS[key],styles["TXHeading"]))
            questions=model["limitations"] or ["No explicit data-gap metadata was returned by the existing analysis."]
            for question in questions: story.append(Paragraph(f"• {_safe(question)}",styles["TXBody"]))
            origin=model["attack_origin"]
            if origin: story.append(Paragraph(f"Attack-origin assessment: {_safe(origin.get('status','unknown'))}. {_safe(origin.get('summary','No summary supplied.'))}",styles["TXBody"]))
        elif key == "technical_appendix":
            story += [Paragraph(SECTION_LABELS[key],styles["TXHeading"]),Paragraph(f"Investigation ID: {_safe(model['case_id'])}",styles["TXBody"]),Paragraph(f"Composition mode: {_safe(model['composition_mode'])}",styles["TXBody"]),Paragraph(f"Persisted event references: {_safe([_event_id(e,i) for i,e in enumerate(model['events'])])}",styles["TXSmall"]),Paragraph(f"Persisted evidence references: {_safe([item['id'] for item in model['evidence']])}",styles["TXSmall"])]
        elif key == "submitted_images":
            story += [Paragraph(SECTION_LABELS[key],styles["TXHeading"]),Paragraph("No supported submitted evidence images are associated with this investigation. Image inclusion is unavailable.",styles["TXBody"])]
        elif key == "methodology":
            story += [Paragraph(SECTION_LABELS[key],styles["TXHeading"]),Paragraph("The report is rendered deterministically from persisted investigation metadata, saved evidence records, timeline events, and the existing TraceX engine result. No provider-backed AI summarization is configured. Evidence approval reflects a reviewer decision, not independent verification. Report generation is read-only and does not alter the case or engine findings.",styles["TXBody"])]
    doc.build(story)
    data=output.getvalue()
    if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-64:]:
        raise ValueError("Generated PDF validation failed.")
    reader = PdfReader(BytesIO(data), strict=True)
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not reader.pages or model["case_id"] not in extracted or model["title"] not in extracted:
        raise ValueError("Generated PDF content validation failed.")
    return data


def _rgb(value: colors.Color) -> RGBColor:
    return RGBColor(round(value.red*255),round(value.green*255),round(value.blue*255))


def _slide(prs: Presentation, title: str) -> Any:
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    background=slide.background.fill; background.solid(); background.fore_color.rgb=RGBColor(16,24,40)
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.13),prs.slide_height); bar.fill.solid();bar.fill.fore_color.rgb=RGBColor(232,106,191);bar.line.fill.background()
    box=slide.shapes.add_textbox(Inches(.55),Inches(.3),Inches(12.1),Inches(.6)); p=box.text_frame.paragraphs[0];p.text=title;p.font.name="Aptos Display";p.font.size=Pt(25);p.font.bold=True;p.font.color.rgb=RGBColor(255,248,253)
    footer=slide.shapes.add_textbox(Inches(.6),Inches(7.12),Inches(12),Inches(.2)); p=footer.text_frame.paragraphs[0];p.text=f"TraceX  ·  {slide.slide_id}";p.font.size=Pt(8);p.font.color.rgb=RGBColor(152,163,179)
    return slide


def _textbox(slide: Any, text: Any, x: float, y: float, w: float, h: float, *, size: int=14, color: RGBColor|None=None, bold: bool=False) -> None:
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.08);tf.margin_right=Inches(.08);tf.margin_top=Inches(.04);tf.margin_bottom=Inches(.04)
    p=tf.paragraphs[0];p.text=_clean(text);p.font.name="Aptos";p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color or RGBColor(220,228,239)


def _add_bullets(prs: Presentation, title: str, lines: list[str]) -> None:
    slide=_slide(prs,title)
    if not lines: lines=["No data is available for this section."]
    for start in range(0,len(lines),6):
        if start:
            slide=_slide(prs,title+" (continued)")
        chunk=lines[start:start+6]
        _textbox(slide,"\n\n".join(f"• {line[:260]}" for line in chunk),.7,1.25,11.8,5.55,size=16)


def generate_powerpoint(model: dict[str, Any], sections: list[str]) -> bytes:
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    cover=_slide(prs,"TraceX Investigation Report")
    _textbox(cover,model["title"],.85,1.8,11.5,1.3,size=30,bold=True,color=RGBColor(255,248,253))
    _textbox(cover,f"Investigation {model['case_id']}\nStatus: {model['status']}\nGenerated: {model['generated_at']}\n\nDeterministic report from persisted TraceX data.",.9,3.35,10.8,2.2,size=16)
    if model["synthetic"]: _textbox(cover,"CONTROLLED SYNTHETIC DATASET",.9,6,10,.4,size=12,bold=True,color=RGBColor(255,209,102))
    if "executive_summary" in sections: _add_bullets(prs,"Executive Summary",[model["root_cause_text"],model["composition_disclosure"]])
    if "evidence_summary" in sections:
        lines=[f"{item['timestamp']} · {item['evidence_type']} · {item['review_status']} · {item['description']}" for item in model["evidence"]]
        _add_bullets(prs,f"Evidence Inventory · {model['evidence_counts']['total']} saved records",lines or ["No persisted evidence records are associated with this investigation."])
    if "timeline" in sections:
        lines=[f"{_timestamp(e)} · {_clean(e.get('event_type',e.get('type')))} · {_clean(e.get('identity'))} · {_event_id(e,i)}" for i,e in enumerate(model["events"])]
        _add_bullets(prs,"Incident Timeline",lines or ["No timeline events are available."])
    if "relationship_map" in sections:
        slide=_slide(prs,"Entity Relationship Map")
        nodes=model["relationship_nodes"][:15];edges=model["relationships"][:20]
        if nodes and edges:
            points={node["id"]:(1.25+(i%5)*2.45,1.5+(i//5)*1.45) for i,node in enumerate(nodes)}
            for edge in edges:
                if edge["source"] in points and edge["target"] in points:
                    x1,y1=points[edge["source"]];x2,y2=points[edge["target"]]
                    line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1+.7),Inches(y1+.25),Inches(x2+.7),Inches(y2+.25));line.line.color.rgb=RGBColor(53,216,255) if edge["basis"]=="Explicit environment relationship" else RGBColor(232,106,191);line.line.width=Pt(1.2)
            for node in nodes:
                x,y=points[node["id"]];shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(1.45),Inches(.5));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(38,33,55);shape.line.color.rgb=RGBColor(53,216,255) if node["type"]!="event" else RGBColor(232,106,191)
                tf=shape.text_frame;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.paragraphs[0].text=node["id"][:22];tf.paragraphs[0].font.size=Pt(8);tf.paragraphs[0].font.color.rgb=RGBColor(245,242,250)
            _textbox(slide,"Cyan: explicit environment relationship. Pink: existing-analysis correlation (inferred).",.7,6.4,11.5,.35,size=10)
        else: _textbox(slide,"No relationship data is available.",.9,1.5,11,.5)
    if "attack_path" in sections:
        lines=[f"{i+1}. {step.get('timestamp') or 'Time not recorded'} · {step.get('event_type',step.get('action','Observed step'))} · {_clean(step.get('identity'))} · evidence {', '.join(map(str,step.get('event_ids',[])))}" for i,step in enumerate(model["attack_path"])]
        _add_bullets(prs,"Potential Attack Path",lines or ["No supported attack path is available in the persisted analysis."])
    if "root_cause" in sections: _add_bullets(prs,"Root-Cause Analysis",[model["root_cause_text"],f"Assessment: {model['root_cause_status']}"]+[f"{f.get('code','Factor')}: {f.get('description','')}" for f in model["root_cause"].get("contributing_factors",[])])
    if "risk_assessment" in sections:
        blast=model["blast_radius"]
        _add_bullets(prs,"Risk Assessment and Potential Impact",[f"Existing engine blast radius: {blast.get('summary',{})}" if blast else "No quantitative risk score or blast-radius output is available.","The report generator does not create or infer risk scores."])
    if "remediation" in sections: _add_bullets(prs,"Recommended Remediation",[f"{r.get('title',r.get('action','Recommendation'))} — {r.get('description',r.get('reason',''))} [{r.get('status','recommended')}]" for r in model["remediations"]] or ["No recommendations are saved for this analysis."])
    if "unresolved_questions" in sections: _add_bullets(prs,"Unresolved Questions and Evidence Gaps",model["limitations"] or ["No explicit gap metadata was returned."])
    if "technical_appendix" in sections: _add_bullets(prs,"Technical Appendix",[f"Investigation: {model['case_id']}",f"Persisted events: {model['event_count']}; included timeline events: {model['events_in_report']}",f"Persisted evidence records: {model['evidence_counts']['total']}",f"Source event IDs: {', '.join(_event_id(e,i) for i,e in enumerate(model['events']))}"])
    if "submitted_images" in sections: _add_bullets(prs,"Submitted Evidence Images",["No supported submitted images are associated with this investigation."])
    if "methodology" in sections: _add_bullets(prs,"Methodology and Limitations",[model["composition_disclosure"],"Approved evidence indicates a review decision only; it is not independent verification.","Report generation is read-only and does not modify investigation data."])
    output=BytesIO();prs.save(output);data=output.getvalue()
    if len(data)>MAX_REPORT_BYTES: raise ValueError("Generated PowerPoint exceeds the configured report size limit.")
    check=Presentation(BytesIO(data))
    if len(check.slides)!=len(prs.slides) or not check.slides: raise ValueError("Generated PowerPoint validation failed.")
    for slide in check.slides:
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width or shape.top+shape.height > prs.slide_height:
                raise ValueError("Generated PowerPoint contains content outside slide boundaries.")
    return data
